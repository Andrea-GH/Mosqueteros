from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from catalogos.models import Categoria, Sucursal, MetodoPago, Cliente, Vendedor, Producto
from ventas.models import Venta, DetalleVenta


class Command(BaseCommand):
    help = "Importa datos desde la hoja BaseDeDatos del archivo DashBoard2021.xlsx"

    def add_arguments(self, parser):
        parser.add_argument(
            '--archivo',
            type=str,
            default='DashBoard2021.xlsx',
            help='Ruta del archivo Excel a importar'
        )

    def normalizar_texto(self, valor, default='Sin dato'):
        if valor is None:
            return default
        texto = str(valor).strip()
        return texto if texto else default

    def decimal_seguro(self, valor, default='0'):
        if valor is None or valor == '':
            return Decimal(default)
        try:
            return Decimal(str(valor))
        except (InvalidOperation, TypeError, ValueError):
            return Decimal(default)

    def generar_codigo_producto(self, nombre_producto):
        base = ''.join(ch for ch in nombre_producto.upper() if ch.isalnum() or ch == ' ')
        base = base.replace(' ', '-')[:30]
        return f"PROD-{base}"

    @transaction.atomic
    def handle(self, *args, **options):
        ruta_archivo = Path(options['archivo'])

        if not ruta_archivo.exists():
            self.stdout.write(self.style.ERROR(f"No se encontró el archivo: {ruta_archivo}"))
            return

        wb = load_workbook(ruta_archivo, data_only=False)
        if 'BaseDeDatos' not in wb.sheetnames:
            self.stdout.write(self.style.ERROR("La hoja 'BaseDeDatos' no existe en el archivo Excel."))
            return

        ws = wb['BaseDeDatos']

        # Encabezados esperados en la fila 2:
        # Documento, Fecha, id_cliente, Cliente, Ciudad, Provincia, Vendedor,
        # Empresa, Forma de pago, Producto, Categoría, Precio, Cantidad, Ventas

        registros = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            documento = row[1]
            fecha = row[2]
            id_cliente = row[3]
            cliente = row[4]
            ciudad = row[5]
            provincia = row[6]
            vendedor = row[7]
            empresa = row[8]
            forma_pago = row[9]
            producto = row[10]
            categoria = row[11]
            precio = row[12]
            cantidad = row[13]
            ventas = row[14]

            if documento is None or cliente is None or producto is None:
                continue

            registros.append({
                'documento': documento,
                'fecha': fecha,
                'id_cliente': id_cliente,
                'cliente': self.normalizar_texto(cliente),
                'ciudad': self.normalizar_texto(ciudad),
                'provincia': self.normalizar_texto(provincia),
                'vendedor': self.normalizar_texto(vendedor),
                'empresa': self.normalizar_texto(empresa),
                'forma_pago': self.normalizar_texto(forma_pago),
                'producto': self.normalizar_texto(producto),
                'categoria': self.normalizar_texto(categoria),
                'precio': self.decimal_seguro(precio),
                'cantidad': int(cantidad) if cantidad is not None else 0,
                'ventas': self.decimal_seguro(ventas if not isinstance(ventas, str) else '0'),
            })

        if not registros:
            self.stdout.write(self.style.WARNING("No se encontraron registros para importar."))
            return

        self.stdout.write("Importando catálogos...")

        # 1. Categorías
        categorias_cache = {}
        for r in registros:
            obj, _ = Categoria.objects.get_or_create(
                nombre=r['categoria'],
                defaults={
                    'descripcion': f"Categoría importada desde Excel",
                    'activo': True
                }
            )
            categorias_cache[r['categoria']] = obj

        # 2. Sucursales
        # Como el Excel trae Ciudad y Provincia, usaremos eso como sucursal real.
        # El país se deja como Ecuador porque las ubicaciones del archivo corresponden a Ecuador.
        sucursales_cache = {}
        for r in registros:
            clave = (r['ciudad'], r['provincia'])
            if clave not in sucursales_cache:
                obj, _ = Sucursal.objects.get_or_create(
                    nombre=r['ciudad'],
                    ciudad=r['ciudad'],
                    estado=r['provincia'],
                    pais='Ecuador',
                    defaults={
                        'direccion': f"Sucursal generada desde BaseDeDatos: {r['ciudad']}, {r['provincia']}",
                        'activo': True
                    }
                )
                sucursales_cache[clave] = obj

        # 3. Métodos de pago
        metodos_cache = {}
        for r in registros:
            obj, _ = MetodoPago.objects.get_or_create(
                nombre=r['forma_pago'],
                defaults={
                    'descripcion': 'Método importado desde Excel',
                    'activo': True
                }
            )
            metodos_cache[r['forma_pago']] = obj

        # 4. Clientes
        clientes_cache = {}
        for r in registros:
            obj, _ = Cliente.objects.get_or_create(
                nombre=r['cliente'],
                defaults={
                    'telefono': '',
                    'correo': '',
                    'ciudad': r['ciudad'],
                    'estado': r['provincia'],
                    'pais': 'Ecuador',
                    'activo': True
                }
            )
            clientes_cache[r['cliente']] = obj

        # 5. Vendedores
        vendedores_cache = {}
        for r in registros:
            clave_sucursal = (r['ciudad'], r['provincia'])
            sucursal = sucursales_cache[clave_sucursal]

            obj, _ = Vendedor.objects.get_or_create(
                nombre=r['vendedor'],
                defaults={
                    'telefono': '',
                    'correo': '',
                    'sucursal': sucursal,
                    'activo': True
                }
            )

            # Si el vendedor ya existía pero no tenía la sucursal correcta, la ajustamos
            if obj.sucursal_id != sucursal.id:
                obj.sucursal = sucursal
                obj.save()

            vendedores_cache[r['vendedor']] = obj

        # 6. Productos
        productos_cache = {}
        for r in registros:
            nombre_producto = r['producto']
            categoria = categorias_cache[r['categoria']]
            codigo = self.generar_codigo_producto(nombre_producto)

            obj, created = Producto.objects.get_or_create(
                codigo=codigo,
                defaults={
                    'nombre': nombre_producto,
                    'categoria': categoria,
                    'precio': r['precio'],
                    'stock': 0,
                    'descripcion': 'Producto importado desde Excel',
                    'activo': True
                }
            )

            if not created:
                obj.nombre = nombre_producto
                obj.categoria = categoria
                obj.precio = r['precio']
                obj.activo = True
                obj.save()

            productos_cache[nombre_producto] = obj

        self.stdout.write("Importando ventas y detalles...")

        # Agrupar por documento para crear una Venta por documento y sus DetalleVenta
        ventas_por_documento = {}
        for r in registros:
            doc = int(r['documento'])
            ventas_por_documento.setdefault(doc, []).append(r)

        for documento, items in ventas_por_documento.items():
            primera = items[0]

            cliente = clientes_cache[primera['cliente']]
            sucursal = sucursales_cache[(primera['ciudad'], primera['provincia'])]
            vendedor = vendedores_cache[primera['vendedor']]
            metodo_pago = metodos_cache[primera['forma_pago']]

            folio = f"DOC-{documento:05d}"

            subtotal_venta = Decimal('0.00')
            for item in items:
                subtotal_venta += item['precio'] * item['cantidad']

            # El Excel no trae impuesto separado, así que se conserva impuesto = 0
            impuesto_venta = Decimal('0.00')
            total_venta = subtotal_venta

            venta, _ = Venta.objects.update_or_create(
                folio=folio,
                defaults={
                    'fecha': primera['fecha'],
                    'cliente': cliente,
                    'sucursal': sucursal,
                    'vendedor': vendedor,
                    'metodo_pago': metodo_pago,
                    'subtotal': subtotal_venta,
                    'impuesto': impuesto_venta,
                    'total': total_venta,
                    'observaciones': f"Empresa de embarque: {primera['empresa']}"
                }
            )

            # Para evitar duplicados si vuelves a ejecutar la importación
            venta.detalles.all().delete()

            for item in items:
                producto = productos_cache[item['producto']]
                subtotal_detalle = item['precio'] * item['cantidad']

                DetalleVenta.objects.create(
                    venta=venta,
                    producto=producto,
                    cantidad=item['cantidad'],
                    precio_unitario=item['precio'],
                    subtotal=subtotal_detalle
                )

        self.stdout.write(self.style.SUCCESS("Importación completada correctamente."))
